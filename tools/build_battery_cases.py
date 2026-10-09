"""
Build the battery cases: one per chemistry family, each with its own roads.

**Copyright notice:** Copyright © 2026 Empa, Matthias Roesslein

    ./.venv/bin/python tools/build_battery_cases.py
    -> data/battery_lfp/input_data/case.xlsx
       data/battery_lmfp/input_data/case.xlsx
       data/battery_nmc_high/input_data/case.xlsx
       data/battery_sodium/input_data/case.xlsx
       data/battery_solid_state/input_data/case.xlsx

Nothing is written over an existing case: the workbook is the INPUT, edited by
hand once it exists (placeholders replaced, numbers revised), so a builder that
rewrote it would throw those edits away. A folder that is already there is
named and left alone; move or remove it yourself to build it again.

WHAT AND WHY
------------
Decided 2026-10-08. The battery is not one case but several, because a recycler
treats a lithium iron phosphate cell, a nickel-manganese-cobalt cell and a
sodium-ion cell differently, and the single `data/battery` case could only carry
one blended set of rates. Each case reads its own chemistry's export
(`chemistries` in its source table) and has its own roads:

    LFP, LMFP   hydrometallurgy  +  direct recycling
    NMC_high    hydrometallurgy  +  pyrometallurgy
    sodium      mechanical treatment (the paper's treatment for sodium-ion),
                and, prepared but not used by default, direct recycling
    solid_state packaging and cables only -- nothing is known of its cell

The pretreatment stays `dismantling`, and the mechanical loss of each road stays
INSIDE that road's coefficients: the report's rates are net of pretreatment, and
a separate mechanical multiplier ahead of them would count the same loss twice.

THREE SETS OF COEFFICIENTS IN ONE WORKBOOK
------------------------------------------
`TCs` and `TCs_improved` carry a `variant` column (src/case_tables.py) and
`run.variants` in src/params_schema.py picks one version of each row:

    tc_set = own    the report's Table 2 for the hydrometallurgical road, and the
                    user's own rows where the user has them; where there is no
                    number of the user's -- the split between roads, pyrometallurgy,
                    direct recycling, the sodium treatment -- the paper's REC
    tc_set = BAU    the paper's business-as-usual numbers throughout
    tc_set = REC    the paper's recovery scenario throughout

    sodium_route = mechanical          every sodium cell goes the paper's way
    sodium_route = mechanical_direct   a share goes to direct recycling instead

Rows that are the same in every version are written once, with `variant` blank.

⚠️ THE ONLY PLACE ANY NUMBER COMES FROM IS ONE OF FOUR SOURCES, and every row
says which in its `source` column:

    the user       Table 2 of documentation/BatteryStudy/Battery_Cell_Recycling_
                   Report_LFP_LMFP_NMC.md, and data/battery's own rows
    the paper      Maisel et al., Waste Management 227 (2027) 115873, and its
                   supplement (documentation/BatteryStudy/*-mmc1.xlsx)
    DERIVED        arithmetic on the paper's numbers, said in the row
    PLACEHOLDER    (Claude, not data) -- a number nobody has; the user replaces it

HOW THE PAPER'S NUMBERS BECOME TWO TABLES
-----------------------------------------
The paper gives, for each route and element, the ESTIMATED overall recovery (the
product of its step coefficients) at 2025 and at 2032 and 2050, with a minimum
and a maximum. The model has two tables, 2030 and 2060, and a straight line
between them. So each of minimum, mode and maximum is fitted by the straight
line through the paper's 2032 and 2050 values and READ OFF at 2030 and at 2060:

    - from 2032 to 2050 the model then reproduces the paper (to within 0.4 points),
      because the paper itself runs straight between its anchor years;
    - 2030 differs where the paper bends before 2032: by 2.1 points on average over
      the 50 rows the cases use, 13.9 at most;
    - after 2050 the paper says nothing, and the line goes on (`continue`).

Putting the paper's 2050 value straight into the 2060 table would be off at 2050 by
2.6 points on average and by up to 11.7 on the steepest rows (the iron, phosphorus and
manganese series of REC); that is why it is not done. Measured on the ESTIMATED rates,
which are what is used -- not on the paper's target series, which are not.

THE PAPER'S NUMBERS ARE USED AS THEY ARE, ANOMALIES INCLUDED -- not repaired.
documentation/BATTERY_ROUTES.md lists the ones found (REC manganese below BAU
manganese in Route 2, for one).

EVERY ROAD IS ONE STEP. The paper's routes are chains (pyrolysis, mechanical
treatment, hydrometallurgy, slag treatment...). Their overall rate per element
is all this model needs, so a road is one process with the overall rate, and the
chain is in the paper.

THE GROUPS. Every resource that leaves a road is a group of two rows --
recovered (or handed on) and lost -- each with its OWN range, drawn on its own
and then conditioned on adding up to 1 by the model's existing rule
(src/sampling.condition_on_sum). The loss row is the mirror of the recovered
row's range, (1 - max, 1 - mode, 1 - min), unless the user's own case already has
a loss row for that same recovered row, in which case the user's is copied.
What a road does not recover at all (oxygen, silicon, the unresolved `rest`) is a
single row, loss = 1.
"""
from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

PAPER = os.path.join(ROOT, 'documentation', 'BatteryStudy',
                     '1-s2.0-S0956053X2600543X-mmc1.xlsx')
REPORT = os.path.join(ROOT, 'documentation', 'BatteryStudy',
                      'Battery_Cell_Recycling_Report_LFP_LMFP_NMC.md')
USER_CASE = os.path.join(ROOT, 'data', 'battery', 'input_data', 'case.xlsx')

EXPORT_DIR = 'data/processed/battery_recovery_draws_by_chemistry'

# ----------------------------------------------------------------------
#  The sources
# ----------------------------------------------------------------------

Triple = tuple[float, float, float]            # (value_min, value, value_max)

PAPER_NAME = 'Maisel et al. 2027 (Waste Management 227, 115873) supplement'
REPORT_NAME = 'the report, Table 2'


def read_paper() -> dict:
    """
    {(scenario, route, element): {year: (min, mode, max)}} -- the paper's
    ESTIMATED overall recovery per route and element.

    Read from the route sheets `LIB_Route<n> BAU|REC`, which hold, per element,
    the step coefficients, their product (the estimate) and the minimum and
    maximum of that product, in blocks for 2025 and for each target year.

    NOT the `BAU_targets` / `REC_targets` sheets. Those are the regulation's
    targets -- lithium 50 % by 2027, 80 % by 2031 -- which the estimates are
    built to meet and sometimes exceed; the REC lithium target in 2028 is 0.50
    where the estimate is 0.874. The estimates are what the paper's own
    coefficient tables contain, and the targets are what they were aimed at.

    Route 4 shares its pyrometallurgical step with Route 1 and is NOT read separately:
    `share()` adds its share to Route 1's and the road then uses Route 1's rates for
    all of it (the aluminium foil is where they differ most: Route 1 0 / 0.2 -> 0.5,
    Route 4 0.84 / 0.95). Pyrometallurgy is 1-4 % of the NMC cells.

    QUIRKS OF THE SHEETS, handled here:
      - an element a route does not treat has estimate 0 but still prints a
        minimum and maximum (the mechanical step's); that is a point mass at 0;
      - a maximum can exceed 1 (1.005) and is clipped;
      - the block header names the year only for the target blocks; the first
        block is the current estimate, 2025.
    """
    book = openpyxl.load_workbook(PAPER, read_only=True, data_only=True)
    found: dict = {}
    for scenario in ('BAU', 'REC'):
        for route in (1, 2, 3, 4, 5):
            rows = [list(r) for r in book[f'LIB_Route{route} {scenario}'].iter_rows(values_only=True)]
            i = 0
            while i < len(rows):
                first = rows[i]
                if not first or first[0] != 'Elements/Materials':
                    i += 1
                    continue
                header = ['' if v is None else str(v) for v in first]
                estimate = next(c for c, v in enumerate(header) if 'stimated recovery rate' in v)
                low_at = max(c for c, v in enumerate(header)
                             if v.strip().lower().startswith('recovery rate min'))
                high_at = max(c for c, v in enumerate(header)
                              if v.strip().lower().startswith('recovery rate max'))
                year_text = re.search(r'Target\s*(20\d\d)', ' '.join(header))
                year = int(year_text.group(1)) if year_text else 2025
                j = i + 2
                while (j < len(rows) and rows[j] and rows[j][0] is not None
                       and rows[j][0] != 'Elements/Materials'
                       and not str(rows[j][0]).startswith('Adjusted')):
                    row = rows[j]
                    if row[estimate] is not None:
                        mode = float(row[estimate])
                        low = float(row[low_at] or 0.0)
                        high = float(row[high_at] or 0.0)
                        found.setdefault((scenario, route, str(row[0]).strip()), {})[year] = \
                            _ordered((low, mode, high)) if mode > 0 else (0.0, 0.0, 0.0)
                    j += 1
                i = j
    return found


def read_route_mix() -> dict:
    """
    {scenario: {year: {route: share}}} from the supplement's Structural_Decomposition
    sheet: how much of all lithium-ion batteries goes down each of the five routes.
    """
    book = openpyxl.load_workbook(PAPER, read_only=True, data_only=True)
    rows = [list(r) for r in book['Structural_Decomposition'].iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(rows) if r and 'OBS_2024' in [str(v) for v in r])
    where = {str(v): c for c, v in enumerate(rows[start]) if v is not None}
    mix: dict = {}
    # ONLY THE FIRST BLOCK. The sheet repeats `Route 1` ... `Route 5` under every
    # table below it -- the efficiency of each route for cobalt, copper, lithium,
    # and so on -- with the same column headings. Reading on past the first block
    # let the last of them (graphite) overwrite the shares, and a direct share of
    # 0.1 % came out as 82 %. It stops at the first row that is not a route.
    for r in rows[start + 1:]:
        name = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ''
        match = re.fullmatch(r'Route ([1-5])', name)
        if not match:
            if mix:
                break
            continue
        route = int(match.group(1))
        for label, column in where.items():
            scenario, _, year = label.partition('_')
            if scenario in ('OBS', 'BAU', 'REC') and r[column] is not None:
                mix.setdefault(scenario, {}).setdefault(int(year), {})[route] = float(r[column])
    return mix


def read_table2() -> dict:
    """
    {(chemistry, element): (triple at 2030, triple at 2060)} from Table 2 of the
    report. Fractions, not per cent. "N/A" rows (an element the chemistry does
    not contain) are left out.
    """
    text = open(REPORT, encoding='utf-8').read()
    block = text[text.index('### Table 2.'):text.index('**Scenario status:**')]
    found = {}
    for line in block.splitlines():
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if len(cells) < 5 or cells[0] not in ('LFP', 'LMFP', 'NMC') or not re.search(r'\d', cells[2]):
            continue
        triples = []
        for cell in cells[2:5]:
            numbers = [float(x) / 100 for x in re.findall(r'\d+(?:\.\d+)?', cell)[:3]]
            triples.append(tuple(numbers))
        found[(cells[0], cells[1])] = (triples[0], triples[2])      # 2030 and 2060; 2050 is not used
    return found


def read_user() -> dict:
    """The user's own case, `data/battery`: its TCs, TCs_improved and processes sheets."""
    return {sheet: pd.read_excel(USER_CASE, sheet_name=sheet, dtype=str).fillna('')
            for sheet in ('TCs', 'TCs_improved', 'processes')}


# ----------------------------------------------------------------------
#  Numbers
# ----------------------------------------------------------------------

def _ordered(triple: Triple) -> Triple:
    """Inside [0, 1] and with min <= mode <= max, to six decimals."""
    low, mode, high = (min(max(float(v), 0.0), 1.0) for v in triple)
    low, high = min(low, mode), max(high, mode)
    return (round(low, 6), round(mode, 6), round(high, 6))


def through(at_2032: float, at_2050: float) -> tuple[float, float]:
    """The straight line through the paper's 2032 and 2050 values, read at 2030 and 2060."""
    slope = (at_2050 - at_2032) / 18.0
    return at_2032 - 2 * slope, at_2050 + 10 * slope


def fitted(at_2032: Triple, at_2050: Triple) -> tuple[Triple, Triple]:
    """Minimum, mode and maximum each through their own two anchors: (2030, 2060)."""
    now, later = zip(*(through(a, b) for a, b in zip(at_2032, at_2050)))
    return _ordered(now), _ordered(later)


def mirror(triple: Triple) -> Triple:
    """The range of what is left over: (1 - max, 1 - mode, 1 - min)."""
    low, mode, high = triple
    return _ordered((1 - high, 1 - mode, 1 - low))


def share_range(share: float, spread: float = 0.5) -> Triple:
    """A share with a range of +- `spread` of itself. The paper gives none."""
    return _ordered((share * (1 - spread), share, share * (1 + spread)))


# ----------------------------------------------------------------------
#  What each chemistry contains
# ----------------------------------------------------------------------
#
# The (element, component) pairs each chemistry's composition holds, from
# RAWCLICVehicleBattery/data/consolidated/<file>_60kWh_400V_pairs.txt and
# `_components.txt`, which the upstream export is built from. Nitrogen and
# fluorine are in those files and are not exported (upstream
# `battery_elements_not_of_interest`), so they are not here.
#
# Written out rather than read, so the case can be built on a machine that does
# not have that project, and so a change in it shows up as a resource without a
# coefficient in stage 01 -- which is the check that says so.

PACK = ('batteryPackSupportFrame', 'batteryPackModuleEnclosuresAndCoolantManifolds',
        'batteryPackThermalConductor', 'batteryPackCables', 'batteryPackCellTerminals')
CELL = ('cathodeActiveMaterial', 'anodeActiveMaterial', 'batteryCellElectrolyte',
        'batteryCellSeparator', 'batteryCellCasing', 'currentCollectorAnode',
        'currentCollectorCathode')
UNITEMISED = 'batteryCellUnitemised'

_LIB_COMMON = [('C', 'anodeActiveMaterial'), ('Li', 'batteryCellElectrolyte'),
               ('Cu', 'batteryPackCables'), ('Al', 'batteryPackCellTerminals'),
               ('Cu', 'batteryPackCellTerminals'),
               ('Al', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
               ('Fe', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
               ('Fe', 'batteryPackSupportFrame'), ('Al', 'batteryPackThermalConductor'),
               ('Li', 'cathodeActiveMaterial'), ('O', 'cathodeActiveMaterial'),
               ('Cu', 'currentCollectorAnode'), ('Al', 'currentCollectorCathode')]
_NA_COMMON = [('C', 'anodeActiveMaterial'), ('Na', 'batteryCellElectrolyte'),
              ('P', 'batteryCellElectrolyte'), ('Cu', 'batteryPackCables'),
              ('Al', 'batteryPackCellTerminals'),
              ('Al', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
              ('Fe', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
              ('Fe', 'batteryPackSupportFrame'), ('Al', 'batteryPackThermalConductor'),
              ('Fe', 'cathodeActiveMaterial'), ('Na', 'cathodeActiveMaterial'),
              ('Al', 'currentCollectorAnode'), ('Al', 'currentCollectorCathode')]

CHEMISTRY = {
    'LFP': dict(components=PACK + CELL,
                pairs=_LIB_COMMON + [('Fe', 'cathodeActiveMaterial'), ('P', 'cathodeActiveMaterial')]),
    'LMFP': dict(components=PACK + CELL,
                 pairs=_LIB_COMMON + [('Fe', 'cathodeActiveMaterial'), ('P', 'cathodeActiveMaterial'),
                                      ('Mn', 'cathodeActiveMaterial')]),
    'NMC_high': dict(components=PACK + CELL,
                     pairs=_LIB_COMMON + [('Si', 'anodeActiveMaterial'),
                                          ('Co', 'cathodeActiveMaterial'),
                                          ('Mn', 'cathodeActiveMaterial'),
                                          ('Ni', 'cathodeActiveMaterial')]),
    'Na_ion_layered': dict(components=PACK + CELL + (UNITEMISED,),
                           pairs=_NA_COMMON + [('O', 'cathodeActiveMaterial'),
                                               ('Cu', 'cathodeActiveMaterial'),
                                               ('Mn', 'cathodeActiveMaterial'),
                                               ('Ni', 'cathodeActiveMaterial')]),
    'Na_ion_prussian_white': dict(components=PACK + CELL + (UNITEMISED,),
                                  pairs=_NA_COMMON + [('C', 'cathodeActiveMaterial')]),
    'solid_state': dict(components=('batteryPackSupportFrame',
                                    'batteryPackModuleEnclosuresAndCoolantManifolds',
                                    'batteryPackThermalConductor', 'batteryPackCables',
                                    'anodeActiveMaterial', 'cathodeActiveMaterial',
                                    'currentCollectorAnode', 'currentCollectorCathode'),
                        pairs=[('Cu', 'batteryPackCables'),
                               ('Al', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
                               ('Fe', 'batteryPackModuleEnclosuresAndCoolantManifolds'),
                               ('Fe', 'batteryPackSupportFrame'), ('Al', 'batteryPackThermalConductor'),
                               ('Cu', 'currentCollectorAnode'), ('Al', 'currentCollectorCathode')]),
}
# The paper's row for each element it names. An element it does not name is mapped
# to the nearest it does, and the row says so (DERIVED).
PAPER_KEY = {('cathodeActiveMaterial', 'Li'): 'Lithium (CAM)',
             ('cathodeActiveMaterial', 'Ni'): 'Nickel (CAM)',
             ('cathodeActiveMaterial', 'Co'): 'Cobalt (CAM)',
             ('cathodeActiveMaterial', 'Mn'): 'Manganese (CAM)',
             ('cathodeActiveMaterial', 'Fe'): 'Iron (CAM)',
             ('cathodeActiveMaterial', 'P'): 'Phosphorous (CAM)',
             ('anodeActiveMaterial', 'C'): 'Graphite (AAM)',
             ('batteryCellElectrolyte', 'Li'): 'Lithium (CAM)',
             ('currentCollectorAnode', 'Cu'): 'Copper (CC)',
             ('currentCollectorAnode', 'Al'): 'Aluminium (CC)',
             ('currentCollectorCathode', 'Al'): 'Aluminium (CC)'}
# For the sodium cells, whose cathodes the paper has no rows for.
GENERIC_CATHODE = 'Lithium (CAM)'

FLOWS = {'hydro': {'Li': 'F_li', 'Ni': 'F_ni', 'Co': 'F_co', 'Mn': 'F_mn', 'Fe': 'F_fe_cell',
                   'P': 'F_p', 'C': 'F_graphite', 'Cu': 'F_cu_cell', 'Al': 'F_al_cell'},
         'pyro': {'Li': 'F_li_pyro', 'Ni': 'F_ni_pyro', 'Co': 'F_co_pyro', 'Mn': 'F_mn_pyro',
                  'Fe': 'F_fe_pyro', 'P': 'F_p_pyro', 'C': 'F_graphite_pyro',
                  'Cu': 'F_cu_pyro', 'Al': 'F_al_pyro'}}

ROADS = {
    'hydro': dict(process='cell_recycling', technology='hydrometallurgical',
                  entry='F_cells_hydro', loss='F_loss_cell'),
    'direct': dict(process='direct_recycling', technology='direct regeneration',
                   entry='F_cells_direct', loss='F_loss_direct'),
    'pyro': dict(process='pyrometallurgy', technology='pyrometallurgical',
                 entry='F_cells_pyro', loss='F_loss_pyro'),
    'mech': dict(process='mechanical_treatment', technology='mechanical',
                 entry='F_cells_mech', loss='F_loss_mech'),
}

UNRESOLVED = 'F_loss_unresolved'
UNRESOLVED_PROCESS = 'unresolved_material'
UNRESOLVED_TECHNOLOGY = 'not itemised upstream'

DEFINITIONAL = 'DEFINITIONAL: not recovered on this road, so all of it is lost. One destination, therefore 1.0.'
UNSPECIFIED = ('Unspecified material -- not recovered, which is what makes every recovery '
               'figure here a lower bound. One destination, therefore 1.0.')
NO_ELEMENTS = ('Unspecified material -- plastics and the like. Upstream names no element in this '
               'component and no road recovers it. Not recovered, which is what makes every recovery '
               'figure here a lower bound. One destination, therefore 1.0.')


@dataclass
class Case:
    folder: str
    chemistries: tuple
    table2: str | None                  # which block of the report's Table 2: LFP, LMFP, NMC
    roads: tuple                        # the roads of the cells, in the order of the split
    note: str = ''
    components: tuple = ()
    pairs: tuple = ()
    sodium: bool = False

    def __post_init__(self):
        components, pairs = [], []
        for name in self.chemistries:
            for c in CHEMISTRY[name]['components']:
                if c not in components:
                    components.append(c)
            for p in CHEMISTRY[name]['pairs']:
                if p not in pairs:
                    pairs.append(p)
        self.components = tuple(components)
        self.pairs = tuple((c, e) for e, c in pairs)         # (component, element)

    @property
    def pack(self):
        return [c for c in self.components if c in PACK]

    @property
    def cell(self):
        return [c for c in self.components if c not in PACK]

    def has(self, component: str, element: str) -> bool:
        return (component, element) in self.pairs


def has_elements(case: 'Case', component: str) -> bool:
    """Whether upstream names at least one element inside this component."""
    return any(c == component for c, _ in case.pairs)


CASES = [
    Case('battery_lfp', ('LFP',), 'LFP', ('hydro', 'direct')),
    Case('battery_lmfp', ('LMFP',), 'LMFP', ('hydro', 'direct')),
    Case('battery_nmc_high', ('NMC_high',), 'NMC', ('hydro', 'pyro')),
    Case('battery_sodium', ('Na_ion_layered', 'Na_ion_prussian_white'), None,
         ('mech', 'direct'), sodium=True),
    Case('battery_solid_state', ('solid_state',), None, ()),
]


# ----------------------------------------------------------------------
#  Rows
# ----------------------------------------------------------------------

COLUMNS = ['Input_FlowID', 'Input_layer', 'Input_layer_key', 'Output_FlowID',
           'TC_target_layer', 'TC_target_key', 'value_min', 'value', 'value_max',
           'process', 'technology', 'source', 'variant']


@dataclass
class Row:
    """One coefficient in both tables: its 2030 numbers and its 2060 numbers."""
    flow_in: str
    layer_in: str
    key_in: str
    flow_out: str
    layer_out: str
    key_out: str
    now: Triple
    later: Triple
    process: str
    technology: str
    source_now: str
    source_later: str
    variant: str = ''

    def as_dict(self, which: str) -> dict:
        low, mode, high = self.now if which == 'now' else self.later
        return {'Input_FlowID': self.flow_in, 'Input_layer': self.layer_in,
                'Input_layer_key': self.key_in, 'Output_FlowID': self.flow_out,
                'TC_target_layer': self.layer_out, 'TC_target_key': self.key_out,
                'value_min': low, 'value': mode, 'value_max': high,
                'process': self.process, 'technology': self.technology,
                'source': self.source_now if which == 'now' else self.source_later,
                'variant': self.variant}


class Sources:
    """The four sources, read once."""

    def __init__(self):
        self.paper = read_paper()
        self.mix = read_route_mix()
        self.table2 = read_table2()
        self.user = read_user()

    # ---- the paper ----------------------------------------------------
    def route(self, scenario: str, route: int, key: str):
        """(2030 triple, 2060 triple, text) for one paper row, by the line through 2032 and 2050."""
        found = self.paper.get((scenario, route, key))
        if not found or 2032 not in found or 2050 not in found:
            raise KeyError(f'the paper has no {scenario} route {route} row {key!r} at 2032 and 2050')
        now, later = fitted(found[2032], found[2050])
        text = (f'{PAPER_NAME}, LIB_Route{route} {scenario}, {key}: estimated overall rate '
                f'{found[2032][1]:.4g} in 2032 and {found[2050][1]:.4g} in 2050 '
                f'(min {found[2032][0]:.4g}/{found[2050][0]:.4g}, max {found[2032][2]:.4g}/'
                f'{found[2050][2]:.4g}); DERIVED: the straight line through them, read at ')
        return now, later, text

    def share(self, scenario: str, kind: str):
        """(2030 share, 2060 share, text) of the cells that take the second road."""
        def at(year):
            r = self.mix[scenario][year]
            if kind == 'direct':                # LFP family: direct against hydro
                return r[5] / (r[2] + r[5])
            return (r[1] + r[4]) / (r[1] + r[2] + r[4])      # NMC: pyro against hydro
        now, later = through(at(2032), at(2050))
        now, later = max(now, 0.0), max(later, 0.0)
        what = ('Route 5 / (Route 2 + Route 5)' if kind == 'direct'
                else '(Route 1 + Route 4) / (Route 1 + Route 2 + Route 4)')
        text = (f'{PAPER_NAME}, Structural_Decomposition {scenario}: {what} = {at(2032):.4g} in 2032 '
                f'and {at(2050):.4g} in 2050, Route 3 (mechanical only, black mass exported) left out; '
                f'DERIVED: the straight line through them, read at ')
        return now, later, text

    # ---- the user's report and case ------------------------------------
    def table2_pair(self, chemistry: str, element: str):
        found = self.table2.get((chemistry, element))
        return found if found else None

    def user_row(self, sheet: str, flow_in: str, key_in: str, key_out: str, flow_out: str):
        """The user's row for one coefficient as (triple, source text), or None."""
        frame = self.user[sheet]
        hit = frame[(frame['Input_FlowID'] == flow_in) & (frame['Input_layer_key'] == key_in)
                    & (frame['TC_target_key'] == key_out) & (frame['Output_FlowID'] == flow_out)]
        if not len(hit):
            return None
        row = hit.iloc[0]
        return ((float(row['value_min']), float(row['value']), float(row['value_max'])),
                str(row['source']))


def close(a: Triple, b: Triple) -> bool:
    return all(abs(x - y) < 1e-9 for x, y in zip(a, b))


# ----------------------------------------------------------------------
#  The rows of each part of a case
# ----------------------------------------------------------------------

def dismantling_rows(case: Case, src: Sources) -> list[Row]:
    """
    The user's dismantling rows for the components this case has.

    A cell component goes to `F_cells`; a pack component to `F_pack_housing`;
    and each is written to BOTH with the other at zero, as the user's case does.
    `batteryCellUnitemised` is new upstream and has no row of the user's, so it
    is dismantled like the other cell components.

    Solid-state sends its cell components to `F_cells_handed`, which is a
    handoff: nothing is known of what is inside, so nothing is claimed about it.
    """
    handed = case.chemistries == ('solid_state',)
    rows = []
    for component in case.components:
        template = component if component != UNITEMISED else 'cathodeActiveMaterial'
        for flow in ('F_pack_housing', 'F_cells', 'F_loss_dismantling'):
            now = src.user_row('TCs', 'F_collected', 'BEV', template, flow)
            later = src.user_row('TCs_improved', 'F_collected', 'BEV', template, flow)
            if component == UNITEMISED:
                text = ('DERIVED: dismantled like the other cell components -- this component is '
                        'new upstream and has no row of the user\'s')
                texts = (text, text)
            else:
                texts = (f'data/battery (the user): {now[1]}', f'data/battery (the user): {later[1]}')
            out = 'F_cells_handed' if (handed and flow == 'F_cells') else flow
            rows.append(Row('F_collected', 'product', 'BEV', out, 'component', component,
                            now[0], later[0], 'dismantling', 'manual', texts[0], texts[1]))
    return rows


def housing_rows(case: Case, src: Sources) -> list[Row]:
    """The user's shredder rows for the pack components this case has -- the same in every version."""
    now, later = src.user['TCs'], src.user['TCs_improved']
    rows = []
    for _, r in now[(now['Input_FlowID'] == 'F_pack_housing')
                    & now['Input_layer_key'].isin(case.pack)].iterrows():
        mate = later[(later['Input_FlowID'] == 'F_pack_housing')
                     & (later['Input_layer_key'] == r['Input_layer_key'])
                     & (later['TC_target_key'] == r['TC_target_key'])
                     & (later['Output_FlowID'] == r['Output_FlowID'])].iloc[0]
        rows.append(Row('F_pack_housing', 'component', r['Input_layer_key'], r['Output_FlowID'],
                        'element', r['TC_target_key'],
                        (float(r['value_min']), float(r['value']), float(r['value_max'])),
                        (float(mate['value_min']), float(mate['value']), float(mate['value_max'])),
                        'general_recycling', 'shredder',
                        f"data/battery (the user): {r['source']}",
                        f"data/battery (the user): {mate['source']}"))
    return rows


def split_rows(case: Case, src: Sources) -> list[Row]:
    """
    The cells' choice of road: for every cell component, a share to each road.

    BOTH ROWS OF EVERY PAIR CARRY A RANGE OF THEIR OWN. They are drawn on their
    own and then made to add up to 1 by the model's existing rule; one is NOT
    computed as 1 minus the other. The minority road's range is +- 50 % of its
    share, which is a placeholder -- the paper gives the shares and no range --
    and the other road's is its mirror, so the same uncertainty is not written
    twice with different bounds.

    Sodium has two versions, selected by `sodium_route`: every cell goes the
    mechanical way (the paper), or half of them go to direct recycling (a
    placeholder: nothing has been published for it).
    """
    if not case.roads:
        return []
    first, second = case.roads
    entry_a, entry_b = ROADS[first]['entry'], ROADS[second]['entry']
    process, technology = 'route_split', 'allocation'
    rows = []

    versions = []                        # (variant tag, share now, share later, text, text later)
    if case.sodium:
        paper = ('Sodium follows the paper: mechanical treatment only, so none of the cells '
                 'goes to direct recycling.')
        versions.append(('sodium_route=mechanical', (0.0, 0.0, 0.0), (0.0, 0.0, 0.0), paper, paper))
        placeholder = ('PLACEHOLDER (Claude, not data): half of the sodium cells to direct recycling, '
                       'range 0.25-0.75. Nothing has been published for direct recycling of sodium-ion '
                       'cells; this is only here so the second road can be switched on.')
        versions.append(('sodium_route=mechanical_direct', (0.25, 0.5, 0.75), (0.25, 0.5, 0.75),
                         placeholder, placeholder))
    else:
        kind = 'direct' if second == 'direct' else 'pyro'
        for scenario, tag in (('BAU', 'tc_set=BAU'), ('REC', 'tc_set=own|REC')):
            now, later, text = src.share(scenario, kind)
            note = ('. Range +-50 % of the share: PLACEHOLDER (Claude, not data) -- the paper gives '
                    'the shares and no range.')
            versions.append((tag, share_range(now), share_range(later),
                             text + '2030' + note, text + '2060' + note))

    # KEYED LIKE THE DISMANTLING ROWS: the resource is the component, whose parent
    # is the product, so the row names `BEV` as its input and the component as its
    # target. Keyed as a component moving to itself, the stranding check does not
    # see it cover the component and reports every cell component as stranded.
    for tag, second_now, second_later, text_now, text_later in versions:
        for component in case.cell:
            if not has_elements(case, component):
                continue
            rows.append(Row('F_cells', 'product', 'BEV', entry_b, 'component', component,
                            second_now, second_later, process, technology, text_now, text_later, tag))
            rows.append(Row('F_cells', 'product', 'BEV', entry_a, 'component', component,
                            mirror(second_now), mirror(second_later), process, technology,
                            text_now, text_later, tag))

    # ⚠️ A COMPONENT UPSTREAM GIVES NO ELEMENTS -- the cell casing, the separator, the
    # sodium cells' unitemised remainder -- IS A LEAF OF THE COMPOSITION: its placeholder
    # material has no children, so no `rest` is derived under it, and a row keyed at
    # `rest` NEVER FIRES. The component's whole mass reaches the road and stops there:
    # no error, nothing in any total. In data/battery's first scenario that is 3.5 % of
    # the cell stream and 2 % of everything collected. These components are not recovered
    # by any road, so they
    # leave here, once, whichever road the rest of the cell takes -- written as a
    # component-level row like every other row of this step, which is the only layer the
    # stranding check asks this flow about.
    #
    # Labelled as what it is, `unresolved_material`, and NOT as the first road's process:
    # the loss figure names each wedge by the process that produces the loss flow, and the
    # structure figure labels each edge with it, so a road's name here would charge
    # hydrometallurgy with mass that never reached it.
    for component in case.cell:
        if not has_elements(case, component):
            rows.append(Row('F_cells', 'product', 'BEV', UNRESOLVED, 'component', component,
                            (1.0, 1.0, 1.0), (1.0, 1.0, 1.0), UNRESOLVED_PROCESS,
                            UNRESOLVED_TECHNOLOGY, NO_ELEMENTS, NO_ELEMENTS))
    return rows


@dataclass
class Item:
    """One resource a road recovers (or hands on): where it is, where it goes, which paper row."""
    component: str
    element: str
    flow: str
    paper_key: str
    kind: str = 'recovered'             # or 'handed'
    note: str = ''                      # how the paper row was chosen, when it is not obvious
    placeholder: bool = False


def plan(case: Case, road: str) -> list[Item]:
    """What a road recovers from this case's cells, and into which flows."""
    items: list[Item] = []
    has = case.has
    cathode = [e for c, e in case.pairs if c == 'cathodeActiveMaterial']
    if road in ('hydro', 'pyro'):
        flows = FLOWS[road]
        for element in ('Li', 'Ni', 'Co', 'Mn', 'Fe', 'P'):
            if element in cathode:
                items.append(Item('cathodeActiveMaterial', element, flows[element],
                                  PAPER_KEY[('cathodeActiveMaterial', element)]))
        if has('anodeActiveMaterial', 'C'):
            items.append(Item('anodeActiveMaterial', 'C', flows['C'], 'Graphite (AAM)'))
        if has('batteryCellElectrolyte', 'Li'):
            items.append(Item('batteryCellElectrolyte', 'Li', flows['Li'], 'Lithium (CAM)',
                              note='the electrolyte lithium is given the cathode lithium row, as the user does'))
        if has('currentCollectorAnode', 'Cu'):
            items.append(Item('currentCollectorAnode', 'Cu', flows['Cu'], 'Copper (CC)'))
        if has('currentCollectorCathode', 'Al'):
            items.append(Item('currentCollectorCathode', 'Al', flows['Al'], 'Aluminium (CC)'))
    elif road == 'direct':
        for element in cathode:
            if element == 'O':
                continue
            key = PAPER_KEY.get(('cathodeActiveMaterial', element))
            placeholder = False
            note = ''
            if element not in ('Li', 'Fe', 'P') or case.sodium:
                key, placeholder = ('Iron (CAM)' if not case.sodium else GENERIC_CATHODE), True
                note = (f'the paper\'s direct route is LFP only and has no row for {element}; the '
                        f'regeneration rate of its iron/lithium/phosphorus rows is carried over')
            items.append(Item('cathodeActiveMaterial', element, 'F_cathode_direct', key,
                              note=note, placeholder=placeholder))
        if has('anodeActiveMaterial', 'C'):
            items.append(Item('anodeActiveMaterial', 'C', 'F_graphite_direct', 'Graphite (AAM)',
                              placeholder=case.sodium,
                              note='regenerated graphite (the paper\'s Route 5)'))
        if has('currentCollectorAnode', 'Cu'):
            items.append(Item('currentCollectorAnode', 'Cu', 'F_cu_direct', 'Copper (CC)'))
        for component in ('currentCollectorAnode', 'currentCollectorCathode'):
            if has(component, 'Al'):
                items.append(Item(component, 'Al', 'F_al_direct', 'Aluminium (CC)',
                                  placeholder=case.sodium))
    elif road == 'mech':
        for element in cathode:
            if element == 'O':
                continue
            key = PAPER_KEY.get(('cathodeActiveMaterial', element))
            note = ''
            if key is None:
                key = GENERIC_CATHODE
                note = (f'the paper names no row for {element} in a cathode; the black-mass rate of '
                        f'its lithium row is used')
            items.append(Item('cathodeActiveMaterial', element, 'F_black_mass', key, 'handed', note))
        if has('anodeActiveMaterial', 'C'):
            items.append(Item('anodeActiveMaterial', 'C', 'F_black_mass', 'Graphite (AAM)', 'handed',
                              'the paper\'s graphite row stands for the anode material of a sodium cell'))
        for component in ('currentCollectorAnode', 'currentCollectorCathode'):
            if has(component, 'Al'):
                items.append(Item(component, 'Al', 'F_al_mech', 'Aluminium (CC)'))
    return items


def paper_version(src: Sources, scenario: str, route: int, item: Item):
    """(recovered now, later, loss now, later, text now, text later) from the paper."""
    now, later, text = src.route(scenario, route, item.paper_key)
    extra = ''
    if item.note:
        extra += f' {item.note[0].upper()}{item.note[1:]}.'
    if item.placeholder:
        extra += (' PLACEHOLDER (Claude, not data): the paper has no number for this resource on '
                  'this road.')
    return now, later, mirror(now), mirror(later), text + '2030.' + extra, text + '2060.' + extra


def own_hydro_version(case: Case, src: Sources, item: Item):
    """
    (recovered now, later, loss now, later, text now, text later) for the user's own hydrometallurgy.

    The cathode elements come from Table 2 of the report, for THIS chemistry.
    Everything else -- graphite, electrolyte lithium, the collectors -- is the
    user's row in data/battery, copied. The loss row is the user's own when the
    user's recovered row IS Table 2's for this chemistry (so the pair is theirs),
    and the mirror of the recovered row otherwise.
    """
    component, element, flow = item.component, item.element, item.flow
    mine_now = src.user_row('TCs', 'F_cells', component, element, flow)
    mine_later = src.user_row('TCs_improved', 'F_cells', component, element, flow)
    loss_now = src.user_row('TCs', 'F_cells', component, element, 'F_loss_cell')
    loss_later = src.user_row('TCs_improved', 'F_cells', component, element, 'F_loss_cell')

    if component == 'cathodeActiveMaterial':
        table = src.table2_pair(case.table2, element)
        if table is None:
            raise KeyError(f'Table 2 has no {case.table2} {element} row')
        now, later = _ordered(table[0]), _ordered(table[1])
        text_now = (f'{REPORT_NAME}, {case.table2} {element}, 2030: '
                    f'min/mode/max {now[0]:.0%}/{now[1]:.0%}/{now[2]:.0%}.')
        text_later = (f'{REPORT_NAME}, {case.table2} {element}, 2060: '
                      f'min/mode/max {later[0]:.0%}/{later[1]:.0%}/{later[2]:.0%}.')
        theirs = (mine_now and mine_later and loss_now and loss_later
                  and close(mine_now[0], now) and close(mine_later[0], later))
        if theirs:
            note = (' The loss row is the user\'s own in data/battery, which has this same '
                    'recovered row.')
            return (now, later, loss_now[0], loss_later[0], text_now + note, text_later + note)
        return (now, later, mirror(now), mirror(later),
                text_now + ' The loss row is the mirror of the range (1 - max, 1 - mode, 1 - min).',
                text_later + ' The loss row is the mirror of the range (1 - max, 1 - mode, 1 - min).')

    if not (mine_now and mine_later and loss_now and loss_later):
        raise KeyError(f'data/battery has no hydrometallurgy rows for {component} {element}')
    return (mine_now[0], mine_later[0], loss_now[0], loss_later[0],
            f'data/battery (the user): {mine_now[1]}', f'data/battery (the user): {mine_later[1]}')


def road_rows(case: Case, src: Sources, road: str) -> list[Row]:
    """
    Every row of one road, in every version.

    What the road recovers is written once per version, each version carrying
    the `variant` tag it belongs to; the resources it does not recover at all
    (oxygen, silicon, the part upstream does not resolve) are written once,
    for every version. A resource whose rate is zero in a version still gets
    its loss row -- 1 -- so the group closes.
    """
    spec = ROADS[road]
    entry, loss_flow = spec['entry'], spec['loss']
    items = plan(case, road)
    rows: list[Row] = []

    def make(item, flow_out, now, later, text_now, text_later, tag):
        return Row(entry, 'component', item.component, flow_out, 'element', item.element,
                   now, later, spec['process'], spec['technology'], text_now, text_later, tag)

    if road == 'hydro':
        scenarios = [('tc_set=own', lambda item: own_hydro_version(case, src, item)),
                     ('tc_set=BAU', lambda item: paper_version(src, 'BAU', 2, item)),
                     ('tc_set=REC', lambda item: paper_version(src, 'REC', 2, item))]
    elif road == 'pyro':
        scenarios = [('tc_set=BAU', lambda item: paper_version(src, 'BAU', 1, item)),
                     ('tc_set=own|REC', lambda item: paper_version(src, 'REC', 1, item))]
    elif road == 'mech':
        scenarios = [('tc_set=BAU', lambda item: paper_version(src, 'BAU', 3, item)),
                     ('tc_set=own|REC', lambda item: paper_version(src, 'REC', 3, item))]
    else:                               # direct: Route 5 is the same in BAU and REC
        for key in {item.paper_key for item in items}:
            a = src.route('BAU', 5, key)[:2]
            b = src.route('REC', 5, key)[:2]
            if not (close(a[0], b[0]) and close(a[1], b[1])):
                raise ValueError(f'Route 5 differs between BAU and REC for {key!r}; the direct road '
                                 f'would need a tag per set')
        scenarios = [('', lambda item: paper_version(src, 'BAU', 5, item))]

    for tag, values in scenarios:
        for item in items:
            now, later, loss_now, loss_later, text_now, text_later = values(item)
            none = not any(now) and not any(later)
            if not none:
                rows.append(make(item, item.flow, now, later, text_now, text_later, tag))
            rows.append(make(item, loss_flow,
                             (1.0, 1.0, 1.0) if none else loss_now,
                             (1.0, 1.0, 1.0) if none else loss_later,
                             text_now if not none else DEFINITIONAL + ' (' + text_now + ')',
                             text_later if not none else DEFINITIONAL + ' (' + text_later + ')', tag))

    # What this road does not recover: every other resource of a cell component that HAS
    # elements, and the `rest` that upstream leaves out of them. A component that has no
    # element at all never reaches a road -- see `split_rows`.
    named = {(item.component, item.element) for item in items}
    for component in case.cell:
        if not has_elements(case, component):
            continue
        for element in [e for c, e in case.pairs if c == component and (c, e) not in named]:
            rows.append(Row(entry, 'component', component, loss_flow, 'element', element,
                            (1.0, 1.0, 1.0), (1.0, 1.0, 1.0), spec['process'], spec['technology'],
                            DEFINITIONAL, DEFINITIONAL))
        rows.append(Row(entry, 'component', component, loss_flow, 'element', 'rest',
                        (1.0, 1.0, 1.0), (1.0, 1.0, 1.0), spec['process'], spec['technology'],
                        UNSPECIFIED, UNSPECIFIED))
    return rows


# ----------------------------------------------------------------------
#  A whole case
# ----------------------------------------------------------------------

def build_case(case: Case, src: Sources) -> dict:
    """The four sheets of one case, as frames."""
    rows: list[Row] = []
    rows += dismantling_rows(case, src)
    rows += split_rows(case, src)
    for road in case.roads:
        rows += road_rows(case, src, road)
    rows += housing_rows(case, src)

    tcs = pd.DataFrame([r.as_dict('now') for r in rows], columns=COLUMNS)
    improved = pd.DataFrame([r.as_dict('later') for r in rows], columns=COLUMNS)
    return {'source': source_frame(case), 'processes': processes_frame(case, rows),
            'TCs': tcs, 'TCs_improved': improved}


def source_frame(case: Case) -> pd.DataFrame:
    return pd.DataFrame([
        ('upstream_dir', EXPORT_DIR),
        ('chemistries', '; '.join(case.chemistries)),
        ('flow', 'collected'), ('product', 'BEV'), ('inflow_flow_id', 'F_collected'),
        ('child_layer', 'element'), ('group_marker', '__component__'),
        ('material_suffix', ''), ('groups', ''), ('draws', 200000),
        # Flat until 2030, a straight line to 2060, and then the same rate on to
        # 2070 -- the user's own settings in data/battery, unchanged.
        ('improvement_start', 2030), ('improvement_end', 2060),
        ('improvement_after_end', 'continue'),
    ], columns=['key', 'value'])


def role_of(flow: str, entries: set) -> str:
    if flow.startswith('F_loss'):
        return 'loss'
    if flow in ('F_black_mass', 'F_cells_handed'):
        return 'handoff'
    if flow in entries or flow in ('F_pack_housing', 'F_cells'):
        return 'intermediate'
    return 'recovered'


def processes_frame(case: Case, rows: list[Row]) -> pd.DataFrame:
    """The flow network, read off the rows so that it cannot disagree with them."""
    entries = {ROADS[road]['entry'] for road in case.roads}
    seen, out = set(), []
    for r in rows:
        key = (r.flow_in, r.flow_out)
        if key in seen:
            continue
        seen.add(key)
        out.append(dict(Input_FlowID=r.flow_in, Output_FlowID=r.flow_out, process=r.process,
                        technology=r.technology, keyed_at=r.layer_out,
                        role=role_of(r.flow_out, entries)))
    return pd.DataFrame(out, columns=['Input_FlowID', 'Output_FlowID', 'process', 'technology',
                                      'keyed_at', 'role'])


def combinations(tcs: pd.DataFrame) -> list[dict]:
    """Every combination of the choices a table offers."""
    from itertools import product

    from src import case_tables
    offered = case_tables.offered(tcs)
    names = list(offered)
    return [dict(zip(names, choice)) for choice in product(*(sorted(offered[n]) for n in names))] \
        or [{}]


RESOURCE = ['Input_FlowID', 'Input_layer', 'Input_layer_key', 'TC_target_layer', 'TC_target_key']
IDENTITY = ['Input_FlowID', 'Input_layer', 'Input_layer_key', 'Output_FlowID',
            'TC_target_layer', 'TC_target_key']


def check_case(name: str, tables: dict) -> list[str]:
    """
    What is wrong with a case's tables, in plain words. Empty means clean.

    For EVERY combination of choices the table offers: the two sheets name the
    same coefficients once each; every number is inside [0, 1] and in order;
    every resource's rows add up to 1 at the mode. A version nobody has selected
    is read the day somebody does, so it is checked now.
    """
    from src import case_tables
    problems = []
    for selection in combinations(tables['TCs']):
        label = ', '.join(f'{k}={v}' for k, v in selection.items()) or 'no choices'
        now = case_tables.select(tables['TCs'], selection)
        later = case_tables.select(tables['TCs_improved'], selection, 'TCs_improved')
        a = now[IDENTITY].astype(str).agg('|'.join, axis=1)
        b = later[IDENTITY].astype(str).agg('|'.join, axis=1)
        if set(a) != set(b):
            problems.append(f'{name} [{label}]: the two sheets do not name the same coefficients '
                            f'({len(set(a) ^ set(b))} differ)')
        for sheet, table, keys in (('TCs', now, a), ('TCs_improved', later, b)):
            if keys.duplicated().any():
                problems.append(f'{name} [{label}] {sheet}: {int(keys.duplicated().sum())} '
                                f'coefficient(s) appear twice')
            low = table['value_min'].astype(float)
            mode = table['value'].astype(float)
            high = table['value_max'].astype(float)
            bad = ((low > mode + 1e-12) | (mode > high + 1e-12) | (low < -1e-12) | (high > 1 + 1e-12))
            if bad.any():
                problems.append(f'{name} [{label}] {sheet}: {int(bad.sum())} row(s) break '
                                f'0 <= min <= mode <= max <= 1')
            totals = mode.groupby([table[c] for c in RESOURCE]).sum()
            off = totals[(totals - 1.0).abs() > 1e-9]
            if len(off):
                problems.append(f'{name} [{label}] {sheet}: {len(off)} resource(s) do not add up to 1, '
                                f'the first is {off.index[0]} = {off.iloc[0]:.6g}')
    return problems


def write_case(case: Case, tables: dict, root: str = ROOT) -> bool:
    """Write one case's workbook. Returns False, writing nothing, if it already exists."""
    from src import case_tables
    folder = os.path.join(root, 'data', case.folder)
    if os.path.exists(case_tables.workbook_path(folder)):
        return False
    for sheet in ('source', 'processes', 'TCs', 'TCs_improved'):
        case_tables.write_sheet(folder, sheet, tables[sheet])
    return True


def build_all(src: Sources | None = None) -> dict:
    src = src or Sources()
    return {case.folder: build_case(case, src) for case in CASES}


def main() -> int:
    src = Sources()
    built = build_all(src)
    failed = False
    for case in CASES:
        tables = built[case.folder]
        problems = check_case(case.folder, tables)
        for problem in problems:
            print(f'  PROBLEM  {problem}', file=sys.stderr)
        if problems:
            failed = True
            continue
        wrote = write_case(case, tables)
        offered = sorted({v for v in tables['TCs']['variant'] if v})
        print(f"{'wrote' if wrote else 'LEFT ALONE (it exists)'}  data/{case.folder}: "
              f"{len(tables['TCs'])} rows in each of TCs and TCs_improved, "
              f"{len(tables['processes'])} processes, {len(offered)} version tags")
    return 1 if failed else 0


if __name__ == '__main__':
    raise SystemExit(main())
